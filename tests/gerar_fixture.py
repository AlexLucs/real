"""Gera tests/fixture.json a partir da aba "Atual" da planilha oficial do BCB
(IBC-Br_todas_publicacoes.xlsx) e calcula os valores esperados com uma
implementação independente das fórmulas, em Python.
Uso: python3 tests/gerar_fixture.py caminho/IBC-Br_todas_publicacoes.xlsx"""
import json, sys, datetime
from decimal import Decimal, ROUND_HALF_UP
import openpyxl

ws = openpyxl.load_workbook(sys.argv[1], data_only=True)["Atual"]
codigos = [c.value for c in ws[10][1:13]]          # linha "SGS"
series = {c: {} for c in codigos}
for row in ws.iter_rows(min_row=11, values_only=True):
    if not isinstance(row[0], datetime.datetime): continue
    for c, v in zip(codigos, row[1:13]):
        if isinstance(v, (int, float)): series[c][f"{row[0].year}-{row[0].month:02d}"] = v

def k(a, m):
    t = a*12 + m-1; return f"{t//12}-{t%12+1:02d}"
def pct(x, y): return round((x/y-1)*100, 10)
def calc(sa, nsa, a, m):
    S = lambda aa, mm: sa[k(aa, mm)]; N = lambda aa, mm: nsa[k(aa, mm)]
    med = lambda f, aa, mm, n: sum(f(aa, mm-n+1+i) for i in range(n))/n
    sm  = lambda f, aa, mm, n: sum(f(aa, mm-n+1+i) for i in range(n))
    it = (m-1)//3*3+1
    tri = [S(a, it+i) if it+i <= m else S(a, m) for i in range(3)]
    ano = [S(a, i+1) if i+1 <= m else S(a, m) for i in range(12)]
    return dict(varMes=pct(S(a,m),S(a,m-1)), varTri=pct(med(S,a,m,3),med(S,a,m-3,3)),
        carregTri=pct(sum(tri)/3,med(S,a,it-1,3)), carregAno=pct(sum(ano)/12,med(S,a-1,12,12)),
        varMes12=pct(N(a,m),N(a-1,m)), varTri4=pct(med(N,a,m,3),med(N,a-1,m,3)),
        varAno=pct(sm(N,a,m,m),sm(N,a-1,m,m)), acum12=pct(sm(N,a,m,12),sm(N,a,m-12,12)))
def fmt(v):
    r = Decimal(repr(v)).quantize(Decimal("0.1"), ROUND_HALF_UP)
    return f"{abs(r) if r == 0 else r}".replace(".", ",")

pares = [(24364,24363),(29602,29601),(29604,29603),(29606,29605),(29610,29609),(29608,29607)]
casos = []
for a, m in [(2005,1),(2020,4),(2024,12),(2025,3),(2026,1),(2026,6),(2026,7)]:
    casos.append({"ano": a, "mes": m, "colunas": [
        {k2: fmt(v) for k2, v in calc(series[s], series[n], a, m).items()} for s, n in pares]})
json.dump({"origem": ws["A1"].value, "series": series, "casos": casos},
          open("tests/fixture.json", "w"), ensure_ascii=False)
print(len(casos), "casos;", casos[-1]["colunas"][0])
